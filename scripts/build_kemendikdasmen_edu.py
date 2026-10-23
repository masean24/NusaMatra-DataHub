#!/usr/bin/env python3
"""
Build verified Kemendikdasmen education indicators at kabupaten/kota level.

Source (official Pusdatin / Portal Data Kemendikdasmen):
- Jumlah ruang kelas SD menurut wilayah, kondisi, dan status, TA 2022/2023
- Jumlah ruang kelas SMP menurut wilayah, kondisi, dan status, TA 2022/2023
- Rasio peserta didik terhadap pendidik SD, TA 2022/2023

Outputs:
- data/kemendikdasmen_edu_kabkota_2023.csv
- data/kemendikdasmen_edu_kabkota_2023.provenance.json

Indicators:
- edu_student_teacher_ratio_sd: direct ratio from the official workbook.
- edu_damaged_classrooms_pct: (rusak sedang + rusak berat) / total ruang kelas
  across SD and SMP, multiplied by 100.
"""

from __future__ import annotations

import csv
import hashlib
import json
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from app.services.crosswalk_service import crosswalk_rows

RAW_DIR = REPO_ROOT / "data" / "raw" / "kemendikdasmen"
OUTPUT_CSV = REPO_ROOT / "data" / "kemendikdasmen_edu_kabkota_2023.csv"
OUTPUT_MANIFEST = REPO_ROOT / "data" / "kemendikdasmen_edu_kabkota_2023.provenance.json"
KAB_GEOJSON = REPO_ROOT / "data" / "boundaries" / "indonesia_kabupaten_low.geojson"

SOURCE_FILES = {
    "ruang_kelas_sd": {
        "filename": "ruang_kelas_sd_2022_2023.xlsx",
        "url": "https://cdn.data-portal.belajar.id/Aset-Dataset/20241105/18573d786866088078bd059aecb2da1069862126ab2d806fd22d1f48e34665eb.xlsx",
        "page": "https://data.kemendikdasmen.go.id/dataset/p/sarana-dan-prasarana/kondisi-ruang-kelas-rusak-berat-sd",
    },
    "ruang_kelas_smp": {
        "filename": "ruang_kelas_smp_2022_2023.xlsx",
        "url": "https://cdn.data-portal.belajar.id/Aset-Dataset/20241105/30b39bf842fad8eb8cb235abd2527bdf6076cf94695e154937861ecf614f0efb.xlsx",
        "page": "https://data.kemendikdasmen.go.id/dataset/p/sarana-dan-prasarana/kondisi-ruang-kelas-rusak-berat-smp",
    },
    "rasio_sd": {
        "filename": "rasio_murid_guru_sd_2022_2023.xlsx",
        "url": "https://cdn.data-portal.belajar.id/Aset-Dataset/20241105/1d05febc22bb5ad9f749ccfa6eee06e9444c08256b3f48977f7ff465c78af9d6.xlsx",
        "page": "https://data.kemendikdasmen.go.id/dataset/p/sarana-dan-prasarana/rasio-murid-guru-sd",
    },
}


def download_if_needed(meta: dict) -> Path:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    path = RAW_DIR / meta["filename"]
    if path.exists() and path.stat().st_size > 0:
        return path
    request = urllib.request.Request(
        meta["url"], headers={"User-Agent": "NusaMatra-DataHub/0.1.0"}
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        path.write_bytes(response.read())
    return path


def load_regions() -> list[dict]:
    data = json.loads(KAB_GEOJSON.read_text(encoding="utf-8"))
    regions = []
    for feature in data["features"]:
        props = feature["properties"]
        regions.append(
            {
                "region_id": str(props["region_id"]),
                "parent_region_id": str(props["parent_region_id"]),
                "bps_code": str(props["bps_code"]),
                "name": props["name"],
                "level": "kota" if str(props.get("type", "")).lower() == "kota" else "kabupaten",
            }
        )
    return regions


def parse_classroom_workbook(path: Path) -> dict[str, dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Data Kab-Kota"]
    result = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        row = list(row) + [None] * (17 - len(row))
        if row[1] is None or row[4] is None:
            continue
        try:
            kode = str(int(float(row[1])))
            bps_code = str(int(float(row[2])))
        except (TypeError, ValueError):
            continue
        number = lambda value: float(value or 0)
        # Columns: Baik (5-7), Rusak Ringan (8-10), Rusak Sedang (11-13),
        # Rusak Berat (14-16), each with Negeri/Swasta/Total.
        total = sum(number(row[i]) for i in (7, 10, 13, 16))
        damaged = number(row[13]) + number(row[16])
        result[kode] = {
            "kode_kemdagri": kode,
            "bps_code": bps_code,
            "region_name": str(row[4]).strip(),
            "total_classrooms": total,
            "damaged_classrooms": damaged,
        }
    wb.close()
    if len(result) != 514:
        raise SystemExit(f"Expected 514 kab/kota rows in {path.name}, got {len(result)}")
    return result


def parse_ratio_workbook(path: Path) -> dict[str, dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Data Kab-Kota"]
    result = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        row = list(row) + [None] * (6 - len(row))
        if row[1] is None or row[4] is None:
            continue
        try:
            kode = str(int(float(row[1])))
            bps_code = str(int(float(row[2])))
        except (TypeError, ValueError):
            continue
        if row[5] is None:
            continue
        result[kode] = {
            "kode_kemdagri": kode,
            "bps_code": bps_code,
            "region_name": str(row[4]).strip(),
            "ratio": float(row[5]),
        }
    wb.close()
    if len(result) != 514:
        raise SystemExit(f"Expected 514 kab/kota rows in {path.name}, got {len(result)}")
    return result


def main() -> None:
    raw_paths = {key: download_if_needed(meta) for key, meta in SOURCE_FILES.items()}
    sd = parse_classroom_workbook(raw_paths["ruang_kelas_sd"])
    smp = parse_classroom_workbook(raw_paths["ruang_kelas_smp"])
    ratio = parse_ratio_workbook(raw_paths["rasio_sd"])

    regions = load_regions()
    source_rows = [
        {
            "kode_kemdagri": kode,
            "region_name": row["region_name"],
            "bps_code": row["bps_code"],
        }
        for kode, row in ratio.items()
    ]
    crosswalk = crosswalk_rows(source_rows, regions)
    if crosswalk.unmatched or crosswalk.ambiguous or len(crosswalk.covered_region_ids) != len(regions):
        raise SystemExit(
            f"Crosswalk incomplete: covered={len(crosswalk.covered_region_ids)}/{len(regions)} "
            f"unmatched={crosswalk.unmatched} ambiguous={crosswalk.ambiguous}"
        )

    kode_by_region: dict[str, str] = {}
    for region_id, rows in crosswalk.rows_by_region.items():
        if len(rows) != 1:
            raise SystemExit(f"Region {region_id} has {len(rows)} source rows")
        kode_by_region[region_id] = rows[0]["kode_kemdagri"]

    region_by_id = {region["region_id"]: region for region in regions}
    output_rows = []
    for region_id in sorted(kode_by_region):
        kode = kode_by_region[region_id]
        if kode not in sd or kode not in smp:
            raise SystemExit(f"Missing classroom data for kode {kode}")
        sd_total = sd[kode]["total_classrooms"]
        smp_total = smp[kode]["total_classrooms"]
        total = sd_total + smp_total
        damaged = sd[kode]["damaged_classrooms"] + smp[kode]["damaged_classrooms"]
        if total <= 0:
            raise SystemExit(f"Zero classroom denominator for kode {kode}")
        output_rows.append(
            {
                "region_id": region_id,
                "region_name": region_by_id[region_id]["name"],
                "kode_kemdagri": kode,
                "bps_code": ratio[kode]["bps_code"],
                "edu_student_teacher_ratio_sd": round(ratio[kode]["ratio"], 4),
                "edu_damaged_classrooms_pct": round(damaged / total * 100.0, 4),
                "sd_total_classrooms": int(sd_total),
                "sd_damaged_classrooms": int(sd[kode]["damaged_classrooms"]),
                "smp_total_classrooms": int(smp_total),
                "smp_damaged_classrooms": int(smp[kode]["damaged_classrooms"]),
                "reference_year": 2023,
                "source_label": "Kemendikdasmen Pusdatin TA 2022/2023 (cut off 30 Nov 2022)",
                "source_url": SOURCE_FILES["rasio_sd"]["page"],
            }
        )

    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(output_rows[0]))
        writer.writeheader()
        writer.writerows(output_rows)

    raw_manifest = {}
    for key, path in raw_paths.items():
        raw_manifest[key] = {
            "path": str(path.relative_to(REPO_ROOT)).replace("\\", "/"),
            "url": SOURCE_FILES[key]["url"],
            "page": SOURCE_FILES[key]["page"],
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "bytes": path.stat().st_size,
        }

    manifest = {
        "retrieved_at": datetime.now(timezone.utc).isoformat(),
        "agency": "Kementerian Pendidikan Dasar dan Menengah (Pusdatin)",
        "year": "Tahun Ajaran 2022/2023",
        "cut_off": "30 November 2022",
        "raw_files": raw_manifest,
        "derived_csv": {
            "path": str(OUTPUT_CSV.relative_to(REPO_ROOT)).replace("\\", "/"),
            "rows": len(output_rows),
            "sha256": hashlib.sha256(OUTPUT_CSV.read_bytes()).hexdigest(),
        },
        "formulas": {
            "edu_student_teacher_ratio_sd": "Direct official ratio (peserta didik / pendidik) per kab/kota.",
            "edu_damaged_classrooms_pct": "(rusak sedang total + rusak berat total) / total ruang kelas, SD + SMP, x 100.",
        },
        "notes": (
            "All rows are official kab/kota workbooks. Region IDs are resolved by name + level + "
            "province using the canonical crosswalk; Kode Kemdagri is retained for audit."
        ),
    }
    OUTPUT_MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"[OK] {len(output_rows)} regions -> {OUTPUT_CSV}")
    print(f"[OK] Provenance -> {OUTPUT_MANIFEST}")


if __name__ == "__main__":
    main()
